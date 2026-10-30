from io import BytesIO
from zipfile import ZipFile

import openpyxl
from fastapi import File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, ConfigDict, Field


class Department(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True, extra="forbid")
    id: int = Field(gt=0, le=9007199254740991)
    collegeName: str | None = Field(default=None, max_length=120)
    facultyName: str | None = Field(default=None, max_length=120)
    majorName: str = Field(min_length=1, max_length=120)


HEADERS = ["학과코드", "단과대학", "학부", "학과"]


def parse_departments(data):
    with ZipFile(BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
            raise ValueError("압축 해제 크기는 30MB 이하여야 합니다.")
    book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False)
    try:
        rows = iter(book.active.values)
        if list(next(rows, ())) != HEADERS:
            raise ValueError("열 순서는 학과코드, 단과대학, 학부, 학과여야 합니다.")
        result, seen = [], set()
        for number, row in enumerate(rows, 2):
            if number > 10001:
                raise ValueError("최대 10000행까지 적용할 수 있습니다.")
            if not any(v is not None and str(v).strip() for v in row):
                continue
            try:
                code = str(row[0] or "").strip()
                if not code.isascii() or not code.isdigit():
                    raise ValueError("학과코드는 양의 정수여야 합니다.")
                values = [str(v).strip() if v is not None else "" for v in row[1:4]]
                item = Department(
                    id=int(code),
                    collegeName=values[0] or None,
                    facultyName=values[1] or None,
                    majorName=values[2],
                )
                if any(value.startswith("=") for value in values):
                    raise ValueError("수식 대신 텍스트를 입력하세요.")
                if item.id in seen:
                    raise ValueError("중복된 학과코드입니다.")
                seen.add(item.id)
                result.append(item.model_dump())
            except (ValueError, IndexError) as exc:
                raise ValueError(f"{number}행: {exc}") from exc
        if not result:
            raise ValueError("적용할 학과가 없습니다.")
        return result
    finally:
        book.close()


def export_departments(items):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "학과 목록"
    sheet.append(HEADERS)
    for item in items:
        sheet.append(
            [
                item["id"],
                item.get("collegeName"),
                item.get("facultyName"),
                item["majorName"],
            ]
        )
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    for col, width in [("A", 14), ("B", 28), ("C", 28), ("D", 32)]:
        sheet.column_dimensions[col].width = width
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def install_departments(app, db):
    @app.post("/api/admin/departments")
    async def add_department(body: Department):
        return (await db("POST", "tt_departments", body.model_dump()))[0]

    @app.put("/api/admin/departments/{department_id}")
    async def update_department(department_id: int, body: Department):
        if department_id != body.id:
            raise HTTPException(422, "학과코드는 변경할 수 없습니다.")
        rows = await db(
            "PATCH",
            "tt_departments",
            body.model_dump(),
            params={"id": f"eq.{department_id}"},
        )
        if not rows:
            raise HTTPException(404, "학과를 찾을 수 없습니다.")
        return rows[0]

    @app.post("/api/admin/departments-import")
    async def import_departments(file: UploadFile = File(...), preview: bool = False):
        data = await file.read(5 * 1024 * 1024 + 1)
        if len(data) > 5 * 1024 * 1024:
            raise HTTPException(413, "5MB 이하의 엑셀만 업로드하세요.")
        try:
            items = parse_departments(data)
        except Exception as exc:
            raise HTTPException(422, f"엑셀 검증 실패: {exc}") from exc
        if preview:
            return items

        await db(
            "POST",
            "tt_departments",
            items,
            params={"on_conflict": "id"},
            prefer="resolution=merge-duplicates,return=representation",
        )
        return await db("GET", "tt_departments", params={"order": "id"})

    @app.get("/api/admin/departments.xlsx")
    async def download_departments():
        items = await db("GET", "tt_departments", params={"order": "id"})
        return Response(
            export_departments(items),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="departments.xlsx"'},
        )
